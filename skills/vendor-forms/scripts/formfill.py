"""Fill vendor forms from your company facts, and remember every answer. No model inside: your agent decides
what each blank means; this script finds the blanks, copies the values in code, and remembers.

  python formfill.py blanks  <form>                         # every blank: id, label, section, remembered key
  python formfill.py keys                                   # the keys a blank can map to, and which are on file
  python formfill.py fill    <form> <out> [--plan plan.json] [--approve k1,k2] [--recipient x@y.com]
  python formfill.py remember-field "<label>" <key> [--section "..."]   # this label means this key, from now on
  python formfill.py remember-fact <key> [--value "..."] [--describe "..."] [--release auto|approve|never]
  python formfill.py import  <filled form> [--plan plan.json]            # learn facts from a form you already filled
  python formfill.py store-doc <name> <file> [--release ...]  |  docs  |  open-doc <name> <tmpdir> [--approved]  |  rotate

Common options: --config config (company-profile.yaml, field-keys.yaml, secrets.yaml.age)  --ops ops (labels.jsonl, audit.jsonl)

Supports .docx (tables and "Label: ____" lines), PDF with form fields, flat PDFs with printed blanks, .xlsx/.xlsm.
Nothing here prints a value: output is labels, sections, key names and statuses. Sensitive values live in an
age-encrypted vault (config/secrets.yaml.age, key at ~/.config/officer-can/age.key) and are typed by the owner at a
hidden prompt, never passed on the command line or through chat.
"""
import argparse, contextlib, datetime, getpass, io, json, os, re, sys

import yaml

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "handlers"))
with contextlib.redirect_stdout(io.StringIO()):     # PyMuPDF can print notices on import; keep stdout pure JSON
    from docx_handler import extract_docx_structure, fill_docx
    from pdf_handler import extract_pdf_structure, fill_pdf
    from xlsx_handler import extract_xlsx_structure, fill_xlsx
    try:
        from flat_pdf import find_blanks, fill_flat
    except Exception:                                # numpy/pymupdf missing: flat PDFs go to manual review
        find_blanks = fill_flat = None

KEY_PATH = os.path.expanduser(os.environ.get("OFFICER_CAN_KEY", "~/.config/officer-can/age.key"))
SKIPS = {"SKIP", "SKIP_OTHER_COMPANY", "SKIP_NOT_APPLICABLE", "NEEDS_HUMAN"}

# ---- where each key's value lives -------------------------------------------------------------------
PROFILE = {  # field key -> company-profile.yaml paths (first non-empty wins); any other key is its own path
    "company.year_established": ["company.year_established", "company.incorporated"],
    "contact.main.name": ["contacts.owner.name"], "contact.main.title": ["contacts.owner.title"],
    "contact.main.email": ["contacts.owner.email"], "contact.main.phone": ["contacts.owner.phone", "company.phone"],
    "contact.purchasing.name": ["contacts.purchasing.name", "contacts.owner.name"],
    "contact.purchasing.email": ["contacts.purchasing.email", "contacts.owner.email"],
    "contact.purchasing.phone": ["contacts.purchasing.phone", "contacts.owner.phone", "company.phone"],
    "contact.ap.name": ["contacts.accounts_payable.name"], "contact.ap.email": ["contacts.accounts_payable.email"],
    "contact.ap.phone": ["contacts.accounts_payable.phone", "company.phone"],
    "contact.qa.name": ["contacts.quality.name"], "contact.qa.email": ["contacts.quality.email", "facts.coa_email"],
    "contact.language": ["contacts.language"],
    "receiving.address": ["facts.receiving.address"], "receiving.instructions": ["facts.receiving.hours"],
    "receiving.contact": ["facts.receiving.contact"], "receiving.company": ["facts.receiving.company"],
    "receiving.street": ["facts.receiving.street"], "receiving.city": ["facts.receiving.city"],
    "receiving.state": ["facts.receiving.state"], "receiving.postal_code": ["facts.receiving.postal_code"],
    "receiving.country": ["facts.receiving.country"], "receiving.shipping_method": ["facts.shipping_method"],
}
SECRET = {"tax.ein": "ein", "tax.resale_id": "state_sales_tax_id", "tax.duns": "duns", "bank.name": "bank_name",
          "bank.address": "bank_address", "bank.phone": "bank_phone",
          "bank.account_number": "bank_account_number", "bank.routing_number": "bank_routing_number"}
DEFAULT_RELEASE = {"ein": "auto", "state_sales_tax_id": "auto", "duns": "auto", "owner_ssn": "never", "signature_image": "never"}

# ---- what's never ours to fill --------------------------------------------------------------------------
OTHER_PARTY = re.compile(r"co-?pack|co-?manufactur|reference|broker|distributor|parent|subsidiar|guarantor", re.I)
SHIP_TO = re.compile(r"ship\s*-?\s*to|deliver|receiving", re.I)
IF_DIFFERENT = re.compile(r"if (different|other than)|if not the same", re.I)
NOT_APPLICABLE = re.compile(r"\((chinese|japanese|korean|thai|arabic|cyrillic|local)[^)]*\)|\bp\.?\s?o\.?\s*box\b|local language", re.I)
NEVER_LABEL = re.compile(r"signature|\bssn\b|social security|date of birth|\bdob\b|credit (limit|line) (requested|request)", re.I)
HEADING_LIKE = re.compile(r"^[\w &/-]{0,40}\b(information|details)$", re.I)
TO_RECEIVING = {"company.legal_name": "receiving.company", "company.address.street": "receiving.street",
                "company.address.city": "receiving.city", "company.address.state": "receiving.state",
                "company.address.postal_code": "receiving.postal_code", "company.address.country": "receiving.country",
                "company.address.full": "receiving.address"}
BLANK_LINE = re.compile(r"^(?P<label>[^:_]{2,80}?)\s*:\s*_*\s*$")

def norm(s):
    return re.sub(r"[^a-z0-9]+", " ", (s or "").lower()).strip()

def dig(d, path):
    for part in path.split("."):
        if isinstance(d, list):
            d = d[0] if d else None
        if not isinstance(d, dict):
            return None
        d = d.get(part)
    return d

def put(d, path, value):
    parts = path.split(".")
    for part in parts[:-1]:
        nxt = d.get(part)
        if isinstance(nxt, list):                    # facts.receiving: write to the first location
            nxt = nxt[0] if nxt else nxt.append({}) or nxt[0]
        elif not isinstance(nxt, dict):
            nxt = d[part] = {}
        d = nxt
    d[parts[-1]] = value

# ---- config, memory, vault --------------------------------------------------------------------------
class Ctx:
    def __init__(self, config, ops):
        self.config, self.ops = config, ops
        self.profile_path = os.path.join(config, "company-profile.yaml")
        self.keys_path = os.path.join(config, "field-keys.yaml")
        self.profile = (yaml.safe_load(open(self.profile_path)) or {}) if os.path.exists(self.profile_path) else {}
        self.keys = (yaml.safe_load(open(self.keys_path)) or {}).get("keys", {}) if os.path.exists(self.keys_path) else {}
        self.vault_path = os.path.join(config, "secrets.yaml.age")
        self._secrets = None

    # remembered answers: ops/labels.jsonl rows {"kind":"form_field","label","section","key"}
    def memory(self):
        p = os.path.join(self.ops, "labels.jsonl"); out = {}
        if os.path.exists(p):
            for line in open(p):
                try:
                    r = json.loads(line)
                except ValueError:
                    continue
                if r.get("kind") == "form_field" and r.get("key"):
                    out[(norm(r.get("label")), norm(r.get("section")))] = r["key"]
        return out

    def recall(self, label, section, mem=None):
        mem = self.memory() if mem is None else mem
        if (norm(label), norm(section)) in mem:
            return mem[(norm(label), norm(section))]
        same = {k for (l, _), k in mem.items() if l == norm(label)}
        return same.pop() if len(same) == 1 else None       # same label anywhere, one consistent answer

    def remember_field(self, label, key, section=""):
        os.makedirs(self.ops, exist_ok=True)
        with open(os.path.join(self.ops, "labels.jsonl"), "a") as f:
            f.write(json.dumps({"kind": "form_field", "label": label, "section": section or "", "key": key,
                                "source": "user", "date": datetime.date.today().isoformat()}) + "\n")

    def audit(self, rows):
        if rows:
            os.makedirs(self.ops, exist_ok=True)
            with open(os.path.join(self.ops, "audit.jsonl"), "a") as f:
                f.writelines(json.dumps(r) + "\n" for r in rows)

    def save_profile(self):
        os.makedirs(self.config, exist_ok=True)
        yaml.safe_dump(self.profile, open(self.profile_path, "w"), sort_keys=False, allow_unicode=True)

    # the vault: {"values": {name: {"value", "release"}}, "documents": {...}}, age-encrypted
    def _age(self):
        try:
            import pyrage
        except ImportError:
            sys.exit("Sensitive values need the age vault: pip install pyrage")
        if not os.path.exists(KEY_PATH):
            os.makedirs(os.path.dirname(KEY_PATH), exist_ok=True)
            ident = pyrage.x25519.Identity.generate()
            fd = os.open(KEY_PATH, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
            with os.fdopen(fd, "w") as f:
                f.write(str(ident) + "\n")
            print(f"Created your vault key at {KEY_PATH}. Save a copy in your password manager.", file=sys.stderr)
        ident = pyrage.x25519.Identity.from_str(open(KEY_PATH).read().strip())
        return pyrage, ident

    def secrets(self):
        if self._secrets is None:
            self._secrets = {"values": {}, "documents": {}}
            if os.path.exists(self.vault_path):
                pyrage, ident = self._age()
                self._secrets = yaml.safe_load(pyrage.decrypt(open(self.vault_path, "rb").read(), [ident])) or self._secrets
                self._secrets.setdefault("values", {})
        return self._secrets

    def save_secrets(self):
        pyrage, ident = self._age()
        data = yaml.safe_dump(self.secrets(), sort_keys=False).encode()
        open(self.vault_path, "wb").write(pyrage.encrypt(data, [ident.to_public()]))

    def has_value(self, key):
        spec = self.keys.get(key, {})
        if spec.get("source") == "secret" or key in SECRET:
            if not os.path.exists(self.vault_path):
                return False
            return bool((self.secrets()["values"].get(SECRET.get(key, key.split(".")[-1])) or {}).get("value"))
        return bool(self.profile_value(key))

    def profile_value(self, key):
        if key == "company.address.full":
            ad = dig(self.profile, "company.address") or {}
            parts = [ad.get(k) for k in ("street", "city", "state", "postal_code")]
            return f"{parts[0]}, {parts[1]}, {parts[2]} {parts[3]}" if all(parts) else ad.get("full")
        if key in ("contact.main.first_name", "contact.main.last_name"):
            name = (dig(self.profile, "contacts.owner.name") or "").split()
            return (name[0] if name else "") if key.endswith("first_name") else " ".join(name[1:])
        if key == "date.today":
            return datetime.date.today().strftime("%m/%d/%Y")
        return next((v for p in PROFILE.get(key, [key]) if (v := dig(self.profile, p))), None)

# ---- finding blanks ------------------------------------------------------------------------------------
def pdf_labels(raw):
    """Printed label and section heading for each PDF text field, read off the page (tooltips are often wrong).
    Words typed inside a box never count as a label, so a filled value is never shown to anyone."""
    try:
        import fitz
    except ImportError:
        return {}
    out, section = {}, ""
    heading = re.compile(r"information|references?|section|details|delivery|shipping|billing|bank|credit|trade|general|contact", re.I)
    for page in fitz.open(stream=raw, filetype="pdf"):
        widgets = [w for w in page.widgets() if w.field_type_string == "Text"]
        boxes = [w.rect for w in page.widgets()]
        words = [x for x in page.get_text("words")
                 if not any(b.x0 - 1 <= (x[0] + x[2]) / 2 <= b.x1 + 1 and b.y0 - 1 <= (x[1] + x[3]) / 2 <= b.y1 + 1 for b in boxes)]
        for w in sorted(widgets, key=lambda w: (round(w.rect.y0), w.rect.x0)):
            r = w.rect; mid = (r.y0 + r.y1) / 2
            prev_x = max([o.rect.x1 for o in widgets if o is not w and abs((o.rect.y0 + o.rect.y1) / 2 - mid) < 6
                          and o.rect.x1 <= r.x0 + 2] or [r.x0 - 220])
            left = sorted([x for x in words if abs((x[1] + x[3]) / 2 - mid) < 6 and x[0] >= prev_x - 2 and x[2] <= r.x0 + 2],
                          key=lambda x: x[0])
            text = re.split(r"_{3,}", " ".join(x[4] for x in left))[-1].strip(" :") if left else ""
            above = " ".join(x[4] for x in sorted([x for x in words if 0 < r.y0 - x[3] < 16 and x[0] < r.x1 and x[2] > r.x0],
                                                   key=lambda x: x[0])).strip(" :")
            if above and "_" not in above and len(above) < 70 and len(above.split()) >= 2 and (heading.search(above) or above.isupper()):
                section = above
            out[w.field_name] = (text, section)
    return out

def read_form(path):
    """-> (fmt, raw, items, notes); items = [{"label", "section", "loc"}]. fmt None = needs a person."""
    raw = open(path, "rb").read(); ext = path.lower().rsplit(".", 1)[-1]
    fmt = {"docx": "docx", "pdf": "pdf", "xlsx": "xlsx", "xlsm": "xlsx"}.get(ext)
    if not fmt:
        return None, raw, [], [f"unsupported file type: .{ext} (use .docx, .pdf, .xlsx)"]
    with contextlib.redirect_stdout(io.StringIO()):
        structure, fillable, notes = {"docx": extract_docx_structure, "pdf": extract_pdf_structure,
                                      "xlsx": lambda b: extract_xlsx_structure(b, keep_vba=(ext == "xlsm"))}[fmt](raw)
    items = []
    if fmt == "docx":
        for t in structure["tables"]:
            section = ""
            for r_idx, cells in enumerate(t["rows"]):
                texts = {c["text"].strip() for c in cells}
                if len(texts) == 1 and next(iter(texts)):
                    section = next(iter(texts)); continue
                for i in range(len(cells) - 1):
                    lab, val = cells[i]["text"].strip(), cells[i + 1]["text"].strip()
                    if lab and not val and not lab.endswith("_") and lab != cells[i + 1]["text"]:
                        items.append({"label": lab.rstrip(":"), "section": section,
                                      "loc": {"type": "table_cell", "table": t["table_index"], "row": r_idx, "col": cells[i + 1]["cell_index"]}})
        for p in structure["paragraphs"]:
            m = BLANK_LINE.match(p["text"])
            if m:
                items.append({"label": m["label"].strip(), "section": "",
                              "loc": {"type": "after_label", "paragraph": p["index"], "label": m["label"].strip() + ":"}})
    elif fmt == "pdf" and fillable:
        printed = pdf_labels(raw)
        for f in structure["fields"]:
            if f.get("type") == "/Tx" and not f.get("current_value"):
                text, section = printed.get(f["name"], ("", ""))
                name = re.sub(r"[_\d]+$", "", f["name"]).strip()
                if text and name and len(name) > len(text) and norm(text).replace(" ", "") in norm(name).replace(" ", ""):
                    text = name
                items.append({"label": (text or name or f.get("hint") or "").strip(), "section": section,
                              "loc": {"type": "field", "name": f["name"]}})
    elif fmt == "pdf":
        if find_blanks is None:
            return None, raw, [], notes + ["flat PDF: pip install pymupdf numpy to fill printed blanks"]
        flat = find_blanks(raw)
        if not flat:
            return None, raw, [], notes + ["no fillable fields or printed blanks found (scanned?): fill by hand"]
        fmt = "flat"; notes = [n for n in notes if "AcroForm" not in n] + [f"flat PDF: {len(flat)} printed blanks; values are written onto the page and read back to check"]
        items = [{"label": b["label"], "section": b["section"], "loc": b["loc"]} for b in flat]
    else:
        import openpyxl
        from openpyxl.utils import column_index_from_string, get_column_letter
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        if not structure["sheets"]:
            # simple sheet with no styled answer boxes: a text label with an empty cell to its right is a blank
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for i in range(len(row) - 1):
                        lab, ans = row[i].value, row[i + 1].value
                        if isinstance(lab, str) and lab.strip() and len(lab) < 80 and ans in (None, "") \
                                and (i == 0 or row[i - 1].value in (None, "")):
                            items.append({"label": lab.strip().rstrip(":"), "section": ws.title,
                                          "loc": {"type": "cell_address", "sheet": ws.title, "cell": row[i + 1].coordinate}})
                            break
        for sh in structure["sheets"]:
            ws = wb[sh["name"]]
            if sh.get("answer_col") and sh.get("answer_col") == sh.get("label_col"):
                nxt = get_column_letter(column_index_from_string(sh["label_col"]) + 1)
                for r in sh.get("rows", []):
                    if r.get("label") and ws[f"{nxt}{r['row']}"].value in (None, ""):
                        items.append({"label": r["label"].strip(), "section": sh["name"],
                                      "loc": {"type": "cell_address", "sheet": sh["name"], "cell": f"{nxt}{r['row']}"}})
                continue
            for r in sh.get("rows", []):
                if r.get("label") and not r.get("current_value") and r.get("answer_cell") \
                        and not str(ws[r["answer_cell"]].value or "").startswith("="):     # never touch formulas
                    items.append({"label": r["label"].strip(), "section": sh["name"],
                                  "loc": {"type": "cell_address", "sheet": sh["name"], "cell": r["answer_cell"]}})
    if not items and fmt != "flat":
        return None, raw, [], notes + ["no blanks found"]
    for i, it in enumerate(items):
        it["id"] = f"b{i + 1}"
    return fmt, raw, items, notes

def auto_skip(label, section):
    if section and IF_DIFFERENT.search(section):
        return "SKIP_NOT_APPLICABLE", "only if different from above"
    if section and OTHER_PARTY.search(section) and not re.search(r"bank", section, re.I):
        return "SKIP_OTHER_COMPANY", "this section is about another company"
    if NEVER_LABEL.search(label):
        return "SKIP", "signatures, SSNs, birth dates and credit requests are always left for you"
    if NOT_APPLICABLE.search(label) or HEADING_LIKE.match(label.strip()):
        return "SKIP_NOT_APPLICABLE", "not a blank for our details"
    return None, ""

# ---- commands ----------------------------------------------------------------------------------------
def cmd_blanks(a, ctx):
    fmt, _, items, notes = read_form(a.form)
    if fmt is None:
        print(json.dumps({"status": "manual_review", "notes": notes}, indent=1)); return
    mem = ctx.memory(); out = []
    for it in items:
        skip, why = auto_skip(it["label"], it["section"])
        row = {"id": it["id"], "label": it["label"], "section": it["section"]}
        if skip:
            row.update(key=skip, why=why)
        elif (k := ctx.recall(it["label"], it["section"], mem)):
            row.update(key=k, why="remembered")
        else:
            row.update(key=None, why="decide: pick a key from `keys`, or SKIP / NEEDS_HUMAN")
        out.append(row)
    todo = sum(1 for r in out if r["key"] is None)
    print(json.dumps({"form": os.path.basename(a.form), "format": fmt, "notes": notes, "blanks": out,
                      "to_decide": todo}, indent=1))

def cmd_keys(a, ctx):
    out = {k: {"describe": v.get("describe", ""), "source": v.get("source"), "on_file": ctx.has_value(k)}
           for k, v in ctx.keys.items() if v.get("source") in ("profile", "secret", "computed")}
    print(json.dumps(out, indent=1))

def cmd_fill(a, ctx):
    fmt, raw, items, notes = read_form(a.form)
    if fmt is None:
        print(json.dumps({"status": "manual_review", "notes": notes}, indent=1)); return
    plan = json.load(open(a.plan)) if a.plan else {}
    approved = {k.strip() for k in (a.approve or "").split(",") if k.strip()}
    mem, ops, report, audit, seen = ctx.memory(), [], [], [], set()
    receiving = (ctx.profile.get("facts") or {}).get("receiving") or []
    has_city = any(norm(it["label"]).replace(" ", "") in ("city", "citystatezip", "citystate") for it in items)
    for it in items:
        label, section, loc = it["label"], it["section"], it["loc"]
        row = {"id": it["id"], "label": label, "section": section}
        planned = plan.get(it["id"], plan.get(label))
        skip, why = auto_skip(label, section)
        if isinstance(planned, dict) and "value" in planned:          # a one-off value typed by the user for this form
            ops.append({"locator": loc, "value": str(planned["value"])})
            report.append(dict(row, key="ONE_OFF", status="filled")); continue
        key = planned or (None if skip else ctx.recall(label, section, mem))
        if skip and not planned:
            report.append(dict(row, key=skip, status="left_blank", why=why)); continue
        if not key:
            report.append(dict(row, key=None, status="undecided")); continue
        if key in SKIPS:
            report.append(dict(row, key=key, status="left_blank" if key != "NEEDS_HUMAN" else "needs_you")); continue
        if key == "company.address.street" and not has_city:
            key = "company.address.full"
        if section and SHIP_TO.search(section) and key in TO_RECEIVING:
            key = TO_RECEIVING[key]                                    # ship-to boxes take the delivery address
        row["key"] = key
        dup = (key, norm(label), norm(section))
        if dup in seen and key != "date.today":
            report.append(dict(row, status="left_blank", why="repeated row: only the first is ours")); continue
        seen.add(dup)
        if key.startswith("receiving.") and len(receiving) > 1:
            report.append(dict(row, status="needs_you", why="several delivery locations on file: which one?")); continue
        if ctx.keys.get(key, {}).get("source") == "secret" or key in SECRET:
            entry = ctx.secrets()["values"].get(SECRET.get(key, key.split(".")[-1])) or {} if os.path.exists(ctx.vault_path) else {}
            rule = entry.get("release") or DEFAULT_RELEASE.get(SECRET.get(key, ""), "approve")
            if not entry.get("value"):
                report.append(dict(row, status="missing"))
            elif rule == "never":
                report.append(dict(row, status="never_release"))
            elif rule == "approve" and key not in approved:
                report.append(dict(row, status="held_for_approval"))
            else:
                ops.append({"locator": loc, "value": str(entry["value"])}); report.append(dict(row, status="filled"))
                audit.append({"time": datetime.datetime.now().isoformat(timespec="seconds"), "form": os.path.basename(a.form),
                              "recipient": a.recipient or "", "key": key, "release": rule})
            continue
        val = ctx.profile_value(key)
        if val:
            ops.append({"locator": loc, "value": str(val)}); report.append(dict(row, status="filled"))
        else:
            report.append(dict(row, status="missing"))
    with contextlib.redirect_stdout(io.StringIO()):
        filled, results = {"docx": fill_docx, "pdf": fill_pdf, "flat": fill_flat,
                           "xlsx": lambda b, o: fill_xlsx(b, o, keep_vba=a.form.lower().endswith(".xlsm"))}[fmt](raw, ops)
    failed = {json.dumps(r.get("locator"), sort_keys=True) for r in results if r.get("status") == "failed"}
    by_loc = {it["id"]: json.dumps(it["loc"], sort_keys=True) for it in items}
    for r in report:
        if r["status"] == "filled" and by_loc[r["id"]] in failed:
            r.update(status="needs_you", why="the value didn't land cleanly on the page: check this blank by hand")
    os.makedirs(os.path.dirname(os.path.abspath(a.out)) or ".", exist_ok=True)
    open(a.out, "wb").write(filled)
    ctx.audit(audit)
    if a.remember:                                   # the plan's decisions become memory for every later form
        for it in items:
            k = plan.get(it["id"], plan.get(it["label"]))
            if isinstance(k, str) and k not in ("NEEDS_HUMAN",) and ctx.recall(it["label"], it["section"], mem) != k:
                ctx.remember_field(it["label"], k, it["section"])
    counts = {}
    for r in report:
        counts[r["status"]] = counts.get(r["status"], 0) + 1
    print(json.dumps({"form": os.path.basename(a.form), "out": a.out, "counts": counts, "notes": notes,
                      "fields": report}, indent=1))

def cmd_remember_field(a, ctx):
    if a.key not in ctx.keys and a.key not in SKIPS:
        sys.exit(f"unknown key {a.key}: run `keys`, or add it with remember-fact --describe")
    ctx.remember_field(a.label, a.key, a.section or "")
    print(json.dumps({"remembered": {"label": a.label, "section": a.section or "", "key": a.key}}))

def cmd_remember_fact(a, ctx):
    if a.key not in ctx.keys:
        if not a.describe:
            sys.exit(f"{a.key} is new: add --describe \"what this is\" so later forms can match it")
        ctx.keys[a.key] = {"source": "secret" if a.secret else "profile", "describe": a.describe}
        keys_doc = (yaml.safe_load(open(ctx.keys_path)) or {}) if os.path.exists(ctx.keys_path) else {}
        keys_doc.setdefault("keys", {})[a.key] = ctx.keys[a.key]
        yaml.safe_dump(keys_doc, open(ctx.keys_path, "w"), sort_keys=False, allow_unicode=True)
    sensitive = a.secret or ctx.keys[a.key].get("source") == "secret" or a.key in SECRET
    if sensitive:
        if a.value is not None:
            sys.exit("Sensitive values are never taken on the command line. Run without --value and type it at the prompt.")
        value = getpass.getpass(f"{a.key} (hidden): ").strip()
        name = SECRET.get(a.key, a.key.split(".")[-1])
        entry = ctx.secrets()["values"].setdefault(name, {})
        entry["value"] = value
        entry["release"] = a.release or entry.get("release") or DEFAULT_RELEASE.get(name, "approve")
        ctx.save_secrets()
        print(json.dumps({"saved": a.key, "where": "vault", "release": entry["release"], "length": len(value)}))
        return
    value = a.value if a.value is not None else input(f"{a.key}: ").strip()
    put(ctx.profile, PROFILE.get(a.key, [a.key])[0], value)
    ctx.save_profile()
    print(json.dumps({"saved": a.key, "where": "company-profile.yaml"}))

def cmd_import(a, ctx):
    """Learn facts from a form you already filled. First run lists the filled slots (labels, never values); the
    plan maps slot ids to keys; then public values go to the profile and sensitive ones to the vault, all in code."""
    fmt = {"docx": "docx", "pdf": "pdf", "xlsx": "xlsx", "xlsm": "xlsx"}.get(a.form.lower().rsplit(".", 1)[-1])
    filled = items_with_values(a.form, fmt) if fmt else []
    if not filled:
        print(json.dumps({"status": "manual_review", "notes": ["no typed values found (flat or scanned form?): enter facts with remember-fact"]}))
        return
    if not a.plan:     # first pass: show the filled slots (labels only) so the agent can decide which key each is
        mem = ctx.memory()
        rows = [dict(it, key=ctx.recall(it["label"], it["section"], mem)) for it in filled]
        print(json.dumps({"form": os.path.basename(a.form), "filled_slots": rows,
                          "next": "write plan.json {id: key} for the ones without a key (SKIP for others), then run import again with --plan"},
                         indent=1))
        return
    plan = json.load(open(a.plan))
    values = filled_values(a.form, fmt)
    saved, skipped = [], []
    for it in filled:
        key = plan.get(it["id"]) or ctx.recall(it["label"], it["section"])
        v = values.get(it["id"])
        if not key or key in SKIPS or not v:
            skipped.append(it["label"]); continue
        if ctx.keys.get(key, {}).get("source") == "secret" or key in SECRET:
            name = SECRET.get(key, key.split(".")[-1])
            e = ctx.secrets()["values"].setdefault(name, {})
            e["value"] = v; e.setdefault("release", DEFAULT_RELEASE.get(name, "approve"))
            saved.append({"key": key, "where": "vault"})
        else:
            put(ctx.profile, PROFILE.get(key, [key])[0], v); saved.append({"key": key, "where": "profile"})
        if plan.get(it["id"]):
            ctx.remember_field(it["label"], key, it["section"])
    ctx.save_profile()
    if any(s["where"] == "vault" for s in saved):
        ctx.save_secrets()
    print(json.dumps({"form": os.path.basename(a.form), "saved": saved, "not_saved": len(skipped)}, indent=1))

DOC_RELEASE = {"w9": "auto", "ein_letter": "auto", "incorporation": "auto", "resale_certificate": "auto",
               "certificate_of_insurance": "approve", "bank_letter": "approve", "financial_statements": "never"}

def cmd_store_doc(a, ctx):
    """Encrypt a document (W-9, EIN letter, insurance certificate...) into config/docs/<name>.age."""
    if not re.match(r"^[a-z0-9_]{1,40}$", a.name):
        sys.exit("name: lowercase letters, digits and _ only (e.g. w9, ein_letter, certificate_of_insurance)")
    pyrage, ident = ctx._age()
    os.makedirs(os.path.join(ctx.config, "docs"), exist_ok=True)
    rel = f"docs/{a.name}.age"
    open(os.path.join(ctx.config, rel), "wb").write(pyrage.encrypt(open(a.file, "rb").read(), [ident.to_public()]))
    docs = ctx.secrets().setdefault("documents", {})
    docs[a.name] = {"file": rel, "filename": os.path.basename(a.file),
                    "release": a.release or (docs.get(a.name) or {}).get("release") or DOC_RELEASE.get(a.name, "approve"),
                    "added": datetime.date.today().isoformat()}
    ctx.save_secrets()
    print(json.dumps({"stored": a.name, "release": docs[a.name]["release"],
                      "next": "delete the plaintext copy you stored from, if it's inside a repo"}))

def cmd_docs(a, ctx):
    docs = ctx.secrets().get("documents", {}) if os.path.exists(ctx.vault_path) else {}
    print(json.dumps({n: {"filename": d.get("filename"), "release": d.get("release"), "added": d.get("added")}
                      for n, d in docs.items()}, indent=1))

def cmd_open_doc(a, ctx):
    """Decrypt one document into a folder OUTSIDE the current repo (to attach it), print the path. Delete it after."""
    dest = os.path.realpath(a.dir); here = os.path.realpath(os.getcwd())
    if dest == here or dest.startswith(here + os.sep):
        sys.exit("decrypt into a temp folder outside the repo, e.g. $(mktemp -d)")
    entry = ctx.secrets().get("documents", {}).get(a.name)
    if not entry:
        sys.exit(f"no document named {a.name}: run `docs`")
    rule = entry.get("release", "approve")
    if rule == "never" or (rule == "approve" and not a.approved):
        sys.exit(f"{a.name} is set to '{rule}': {'never shared' if rule == 'never' else 'ask the owner, then pass --approved'}")
    pyrage, ident = ctx._age()
    os.makedirs(dest, exist_ok=True)
    out = os.path.join(dest, entry.get("filename") or f"{a.name}.pdf")
    open(out, "wb").write(pyrage.decrypt(open(os.path.join(ctx.config, entry["file"]), "rb").read(), [ident]))
    ctx.audit([{"time": datetime.datetime.now().isoformat(timespec="seconds"), "form": "document",
                "recipient": a.recipient or "", "key": f"document.{a.name}", "release": rule}])
    print(json.dumps({"path": out}))

def cmd_rotate(a, ctx):
    """New key; re-encrypt the vault and every document; keep the old key as age.key.old until the owner confirms."""
    pyrage, old = ctx._age()
    data = ctx.secrets()
    docs = {n: pyrage.decrypt(open(os.path.join(ctx.config, d["file"]), "rb").read(), [old])
            for n, d in data.get("documents", {}).items()}
    new = pyrage.x25519.Identity.generate()
    for n, blob in docs.items():
        open(os.path.join(ctx.config, data["documents"][n]["file"]), "wb").write(pyrage.encrypt(blob, [new.to_public()]))
    open(ctx.vault_path, "wb").write(pyrage.encrypt(yaml.safe_dump(data, sort_keys=False).encode(), [new.to_public()]))
    os.replace(KEY_PATH, KEY_PATH + ".old")
    fd = os.open(KEY_PATH, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    with os.fdopen(fd, "w") as f:
        f.write(str(new) + "\n")
    print(json.dumps({"rotated": True, "documents": len(docs),
                      "next": f"save the new key in your password manager; delete {KEY_PATH}.old once everything opens"}))

def items_with_values(path, fmt):
    """Label/section/id for every answer slot that HAS a value (for importing a filled form)."""
    raw = open(path, "rb").read(); out = []
    if fmt == "docx":
        from docx import Document
        doc = Document(io.BytesIO(raw))
        for ti, t in enumerate(doc.tables):
            section = ""
            for ri, row in enumerate(t.rows):
                cells = [c.text.strip() for c in row.cells]
                if len(set(cells)) == 1 and cells[0]:
                    section = cells[0]; continue
                for i in range(len(cells) - 1):
                    if cells[i] and cells[i + 1] and cells[i] != cells[i + 1]:
                        out.append({"id": f"t{ti}r{ri}c{i + 1}", "label": cells[i].rstrip(":"), "section": section})
    elif fmt == "pdf":
        from pypdf import PdfReader
        printed = pdf_labels(raw)
        for name, f in (PdfReader(io.BytesIO(raw)).get_fields() or {}).items():
            if f.get("/FT") == "/Tx" and f.get("/V"):
                text, section = printed.get(name, ("", ""))
                out.append({"id": f"f:{name}", "label": text or re.sub(r"[_\d]+$", "", name), "section": section})
    else:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for i in range(len(row) - 1):
                    lab, val = row[i].value, row[i + 1].value
                    if isinstance(lab, str) and lab.strip() and val not in (None, "") and not str(val).startswith("="):
                        out.append({"id": f"{ws.title}!{row[i + 1].coordinate}", "label": lab.strip().rstrip(":"), "section": ws.title})
    return out

def filled_values(path, fmt):
    raw = open(path, "rb").read(); out = {}
    if fmt == "docx":
        from docx import Document
        doc = Document(io.BytesIO(raw))
        for ti, t in enumerate(doc.tables):
            for ri, row in enumerate(t.rows):
                for ci, c in enumerate(row.cells):
                    out[f"t{ti}r{ri}c{ci}"] = c.text.strip()
    elif fmt == "pdf":
        from pypdf import PdfReader
        for name, f in (PdfReader(io.BytesIO(raw)).get_fields() or {}).items():
            out[f"f:{name}"] = str(f.get("/V") or "").strip()
    else:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if c.value not in (None, ""):
                        out[f"{ws.title}!{c.coordinate}"] = str(c.value).strip()
    return out

def main():
    ap = argparse.ArgumentParser(description="Fill vendor forms from your facts, and remember every answer.")
    ap.add_argument("--config", default="config"); ap.add_argument("--ops", default="ops")
    sub = ap.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("blanks"); b.add_argument("form")
    sub.add_parser("keys")
    f = sub.add_parser("fill"); f.add_argument("form"); f.add_argument("out")
    f.add_argument("--plan"); f.add_argument("--approve", default=""); f.add_argument("--recipient", default="")
    f.add_argument("--no-remember", dest="remember", action="store_false")
    rf = sub.add_parser("remember-field"); rf.add_argument("label"); rf.add_argument("key"); rf.add_argument("--section", default="")
    rv = sub.add_parser("remember-fact"); rv.add_argument("key"); rv.add_argument("--value")
    rv.add_argument("--describe"); rv.add_argument("--secret", action="store_true")
    rv.add_argument("--release", choices=["auto", "approve", "never"])
    im = sub.add_parser("import"); im.add_argument("form"); im.add_argument("--plan")
    sd = sub.add_parser("store-doc"); sd.add_argument("name"); sd.add_argument("file")
    sd.add_argument("--release", choices=["auto", "approve", "never"])
    sub.add_parser("docs")
    od = sub.add_parser("open-doc"); od.add_argument("name"); od.add_argument("dir")
    od.add_argument("--approved", action="store_true"); od.add_argument("--recipient", default="")
    sub.add_parser("rotate")
    a = ap.parse_args()
    ctx = Ctx(a.config, a.ops)
    {"blanks": cmd_blanks, "keys": cmd_keys, "fill": cmd_fill, "remember-field": cmd_remember_field,
     "remember-fact": cmd_remember_fact, "import": cmd_import, "store-doc": cmd_store_doc, "docs": cmd_docs,
     "open-doc": cmd_open_doc, "rotate": cmd_rotate}[a.cmd](a, ctx)

if __name__ == "__main__":
    main()
