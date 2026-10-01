"""
.docx extraction and filling using python-docx.

EXTRACTION format:
{
    "paragraphs": [
        {"index": 0, "text": "Vendor Information Form"},
        {"index": 1, "text": "Legal Name: ____________________"},
        ...
    ],
    "tables": [
        {
            "table_index": 0,
            "rows": [
                [{"cell_index": 0, "text": "Legal Name"}, {"cell_index": 1, "text": ""}],
                [{"cell_index": 0, "text": "EIN"},        {"cell_index": 1, "text": ""}],
            ]
        }
    ]
}

FILL OPERATION locator shapes:

  Set the text of a specific table cell (most common for vendor forms):
    {"type": "table_cell", "table": 0, "row": 1, "col": 1}

  Replace a specific substring anywhere it appears in the document:
    {"type": "replace_text", "find": "[Insert Company Name]"}

  Insert text right after a label in a paragraph (handles "Legal Name: _____"):
    {"type": "after_label", "paragraph": 5, "label": "Legal Name:"}

  Skip — used to record fields intentionally left blank (bank info, signature):
    {"type": "skip", "field": "Bank account number"}
"""

from __future__ import annotations

import io
import logging
from typing import Any

from docx import Document
from docx.document import Document as _DocxDocument

log = logging.getLogger("form-fill.docx")


# ---- extraction --------------------------------------------------------------

def extract_docx_structure(raw: bytes) -> tuple[dict, bool, list[str]]:
    """Walk the doc and emit a positional map used to build fill ops."""
    doc = Document(io.BytesIO(raw))

    paragraphs: list[dict] = []
    for idx, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text:  # skip empty paragraphs — they're noise
            paragraphs.append({"index": idx, "text": text})

    tables: list[dict] = []
    for t_idx, table in enumerate(doc.tables):
        rows: list[list[dict]] = []
        for r_idx, row in enumerate(table.rows):
            cells: list[dict] = []
            for c_idx, cell in enumerate(row.cells):
                cells.append({"cell_index": c_idx, "text": cell.text.strip()})
            rows.append(cells)
        tables.append({"table_index": t_idx, "rows": rows})

    structure = {"paragraphs": paragraphs, "tables": tables}

    notes: list[str] = []
    fillable = bool(paragraphs or tables)
    if not fillable:
        notes.append("Document appears empty — no paragraphs or tables found.")
    if len(tables) == 0 and len(paragraphs) > 50:
        notes.append("No tables and >50 paragraphs — likely a prose-heavy form, fill quality may be lower.")

    return structure, fillable, notes


# ---- filling -----------------------------------------------------------------

def fill_docx(raw: bytes, operations: list[dict]) -> tuple[bytes, list[dict]]:
    """Apply each fill operation in order. Return (filled_bytes, per-op report)."""
    doc = Document(io.BytesIO(raw))
    report: list[dict] = []

    for op in operations:
        locator = op.get("locator") or {}
        value = op.get("value", "")
        note = op.get("note")
        loc_type = locator.get("type")

        try:
            if loc_type == "skip":
                report.append({
                    "locator": locator,
                    "value": value,
                    "status": "skipped",
                    "detail": note or "intentionally skipped (e.g., bank info, signature)",
                })

            elif loc_type == "table_cell":
                _fill_table_cell(doc, locator, value)
                report.append({"locator": locator, "value": value, "status": "filled", "detail": note})

            elif loc_type == "replace_text":
                ok = _replace_text(doc, locator.get("find", ""), value)
                report.append({
                    "locator": locator,
                    "value": value,
                    "status": "filled" if ok else "failed",
                    "detail": note if ok else f"substring not found: {locator.get('find', '')!r}",
                })

            elif loc_type == "after_label":
                ok = _after_label(doc, locator, value)
                report.append({
                    "locator": locator,
                    "value": value,
                    "status": "filled" if ok else "failed",
                    "detail": note if ok else f"label not found in paragraph {locator.get('paragraph')}",
                })

            else:
                report.append({
                    "locator": locator,
                    "value": value,
                    "status": "failed",
                    "detail": f"unknown locator type: {loc_type!r}",
                })

        except Exception as e:
            log.exception("op failed: %s", op)
            report.append({"locator": locator, "value": value, "status": "failed", "detail": str(e)})

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue(), report


# ---- low-level helpers -------------------------------------------------------

def _fill_table_cell(doc: _DocxDocument, locator: dict, value: str) -> None:
    """Set a table cell's text. Replaces all content in the cell."""
    t_idx = int(locator["table"])
    r_idx = int(locator["row"])
    c_idx = int(locator["col"])

    cell = doc.tables[t_idx].rows[r_idx].cells[c_idx]
    # Clear existing content but keep the first paragraph (cells must have ≥1 paragraph)
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    first_p = cell.paragraphs[0]
    for run in first_p.runs:
        run.text = ""
    if first_p.runs:
        first_p.runs[0].text = value
    else:
        first_p.add_run(value)


def _replace_text(doc: _DocxDocument, find: str, value: str) -> bool:
    """Replace every occurrence of `find` with `value` across paragraphs and table cells.

    Note: python-docx splits text across runs, so we operate on paragraph.text via a
    rebuild. This loses inline formatting in matched paragraphs, which is acceptable
    for form fields where we're filling blanks.
    """
    if not find:
        return False

    found = False

    def _replace_in_paragraph(p) -> None:
        nonlocal found
        if find in p.text:
            new_text = p.text.replace(find, value)
            # Clear all runs, write new text in a single run
            for run in p.runs:
                run.text = ""
            if p.runs:
                p.runs[0].text = new_text
            else:
                p.add_run(new_text)
            found = True

    for p in doc.paragraphs:
        _replace_in_paragraph(p)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _replace_in_paragraph(p)

    return found


def _after_label(doc: _DocxDocument, locator: dict, value: str) -> bool:
    """Insert `value` immediately after `label` in the specified paragraph.

    Handles common patterns like "Legal Name: ______" — we replace the label + any
    trailing whitespace/underscores with "Legal Name: EXAMPLE FOODS INC.".
    """
    p_idx = int(locator["paragraph"])
    label = locator.get("label", "")

    if p_idx >= len(doc.paragraphs):
        return False
    p = doc.paragraphs[p_idx]
    if label not in p.text:
        return False

    # Replace the label + everything after it on that paragraph with "label value"
    before, _, after = p.text.partition(label)
    # Strip leading underscores/spaces/colons/tabs from `after` (the blank placeholder)
    cleaned_after = after.lstrip(" _:\t")
    new_text = f"{before}{label} {value}".rstrip()
    if cleaned_after:
        new_text = f"{new_text}  {cleaned_after}"

    for run in p.runs:
        run.text = ""
    if p.runs:
        p.runs[0].text = new_text
    else:
        p.add_run(new_text)
    return True
