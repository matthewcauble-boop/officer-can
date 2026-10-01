"""End-to-end check of formfill.py on made-up forms for a made-up company. Run: python test_formfill.py
Builds a .docx, .xlsx and fillable PDF, then: blanks -> plan -> fill -> remember -> second form fills from memory;
release rules (auto / approve / never); sensitive values only via the hidden prompt; nothing ever prints a value."""
import io, json, os, shutil, subprocess, sys, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATES = os.path.join(HERE, "..", "..", "..", "templates", "config")
T = tempfile.mkdtemp()
os.environ["OFFICER_CAN_KEY"] = os.path.join(T, "age.key")
CONFIG, OPS = os.path.join(T, "config"), os.path.join(T, "ops")
os.makedirs(CONFIG)
for f in ("field-keys.yaml", "company-profile.example.yaml"):
    shutil.copy(os.path.join(TEMPLATES, f), os.path.join(CONFIG, f.replace(".example", "")))
EIN, ACCT = "12-3456789", "000123456789"          # fake values; the test asserts they never appear in output

def run(*args, stdin=None):
    p = subprocess.run([sys.executable, os.path.join(HERE, "formfill.py"), "--config", CONFIG, "--ops", OPS, *args],
                       capture_output=True, text=True, input=stdin, env=dict(os.environ))
    assert p.returncode == 0, p.stderr
    for secret in (EIN, ACCT):
        assert secret not in p.stdout and secret not in p.stderr, f"a sensitive value was printed by {args[0]}"
    return json.loads(p.stdout)

# ---- sensitive values go in through the hidden prompt only ----
sys.path.insert(0, HERE)
import formfill, getpass                                       # noqa: E402
for key, val, rel in (("tax.ein", EIN, None), ("bank.account_number", ACCT, None)):
    getpass.getpass = lambda prompt="", v=val: v
    sys.argv = ["formfill.py", "--config", CONFIG, "--ops", OPS, "remember-fact", key] + (["--release", rel] if rel else [])
    out = io.StringIO(); sys.stdout = out; formfill.main(); sys.stdout = sys.__stdout__
    assert EIN not in out.getvalue() and ACCT not in out.getvalue()
assert os.path.exists(os.path.join(CONFIG, "secrets.yaml.age"))
raw_vault = open(os.path.join(CONFIG, "secrets.yaml.age"), "rb").read()
assert EIN.encode() not in raw_vault and ACCT.encode() not in raw_vault, "vault is not encrypted"
p = subprocess.run([sys.executable, os.path.join(HERE, "formfill.py"), "--config", CONFIG, "--ops", OPS,
                    "remember-fact", "tax.ein", "--value", EIN], capture_output=True, text=True)
assert p.returncode != 0, "a sensitive value was accepted on the command line"

# a brand-new fact, remembered for every later form
run("remember-fact", "company.fax", "--value", "(555) 555-0199", "--describe", "Fax number")

# ---- a Word form ----
import docx                                                      # noqa: E402
d = docx.Document(); t = d.add_table(rows=10, cols=2)
rows = [("Company Information", "Company Information"), ("Company Legal Name", ""), ("Main Phone", ""),
        ("Federal Tax ID (EIN)", ""), ("Fax", ""), ("Bank Information", "Bank Information"), ("Account Number", ""),
        ("Trade References", "Trade References"), ("Company Name", ""), ("Authorized Signature", "")]
for i, (a, b) in enumerate(rows):
    t.cell(i, 0).text = a; t.cell(i, 1).text = b
    if a == b:
        t.cell(i, 0).merge(t.cell(i, 1))
d.add_paragraph("Website: ________")
form1 = os.path.join(T, "acme-new-customer.docx"); d.save(form1)

b = run("blanks", form1)
by_label = {x["label"]: x for x in b["blanks"]}
assert by_label["Company Name"]["key"] == "SKIP_OTHER_COMPANY", by_label["Company Name"]
assert by_label["Authorized Signature"]["key"].startswith("SKIP")
undecided = [x for x in b["blanks"] if x["key"] is None]
# the agent decides the undecided ones (here: by hand, as a model would)
choose = {"Company Legal Name": "company.legal_name", "Main Phone": "company.phone", "Federal Tax ID (EIN)": "tax.ein",
          "Fax": "company.fax", "Account Number": "bank.account_number", "Website": "company.website"}
plan = {x["id"]: choose[x["label"]] for x in undecided}
json.dump(plan, open(os.path.join(T, "plan.json"), "w"))
out1 = os.path.join(T, "filled", "acme.docx")
r = run("fill", form1, out1, "--plan", os.path.join(T, "plan.json"), "--recipient", "ap@acme.example")
st = {f["label"]: f["status"] for f in r["fields"]}
assert st["Company Legal Name"] == "filled" and st["Federal Tax ID (EIN)"] == "filled" and st["Fax"] == "filled"
assert st["Account Number"] == "held_for_approval", st          # bank details wait for the owner
assert st["Company Name"] == "left_blank" and st["Authorized Signature"] == "left_blank"
text = "\n".join(c.text for row in docx.Document(out1).tables[0].rows for c in row.cells)
assert "Example Foods Inc." in text and EIN in text and ACCT not in text
assert "example.com" in "\n".join(p.text for p in docx.Document(out1).paragraphs)
r = run("fill", form1, out1, "--plan", os.path.join(T, "plan.json"), "--approve", "bank.account_number")
assert {f["label"]: f["status"] for f in r["fields"]}["Account Number"] == "filled"
audit = [json.loads(l) for l in open(os.path.join(OPS, "audit.jsonl"))]
assert {a["key"] for a in audit} >= {"tax.ein", "bank.account_number"} and all("value" not in a for a in audit)

# ---- an Excel form from someone else: every label it shares with the Word form fills from memory ----
import openpyxl                                                 # noqa: E402
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Vendor Setup"
for i, lab in enumerate(["Company Legal Name", "Main Phone", "Federal Tax ID (EIN)", "Fax", "Total"], start=1):
    ws[f"A{i}"] = lab
ws["B5"] = "=1+1"                                                # a formula cell is never overwritten
form2 = os.path.join(T, "globex-setup.xlsx"); wb.save(form2)
b2 = run("blanks", form2)
assert all(x["key"] for x in b2["blanks"] if x["label"] != "Total"), b2["blanks"]
assert {x["why"] for x in b2["blanks"] if x["label"] != "Total"} == {"remembered"}
out2 = os.path.join(T, "filled", "globex.xlsx")
r2 = run("fill", form2, out2)
ws2 = openpyxl.load_workbook(out2).active
assert ws2["B1"].value == "Example Foods Inc." and ws2["B3"].value == EIN and ws2["B5"].value == "=1+1"

# ---- a fillable PDF with printed labels ----
import pymupdf                                                  # noqa: E402
doc = pymupdf.open(); page = doc.new_page()
for i, (lab, name) in enumerate([("Company Legal Name", "f1"), ("Main Phone", "f2"), ("Shipping Contact", "f3")]):
    y = 100 + i * 40
    page.insert_text((50, y + 14), lab + ":", fontsize=11)
    w = pymupdf.Widget(); w.field_type = pymupdf.PDF_WIDGET_TYPE_TEXT; w.field_name = name
    w.rect = pymupdf.Rect(220, y, 450, y + 20); page.add_widget(w)
form3 = os.path.join(T, "initech.pdf"); doc.save(form3)
b3 = run("blanks", form3)
labels = {x["label"]: x for x in b3["blanks"]}
assert labels["Company Legal Name"]["key"] == "company.legal_name" and labels["Shipping Contact"]["key"] is None
out3 = os.path.join(T, "filled", "initech.pdf")
r3 = run("fill", form3, out3)
assert {f["label"]: f["status"] for f in r3["fields"]}["Shipping Contact"] == "undecided"
assert "Example Foods Inc." in pymupdf.open(out3)[0].get_text() or \
    any(w.field_value == "Example Foods Inc." for w in pymupdf.open(out3)[0].widgets())

# ---- import facts from a form someone already filled ----
d = docx.Document(); t = d.add_table(rows=2, cols=2)
t.cell(0, 0).text = "Year Established"; t.cell(0, 1).text = "2019"
t.cell(1, 0).text = "Number of Employees"; t.cell(1, 1).text = "12"
old = os.path.join(T, "old-form.docx"); d.save(old)
lst = run("import", old)
ids = {x["label"]: x["id"] for x in lst["filled_slots"]}
json.dump({ids["Year Established"]: "company.year_established", ids["Number of Employees"]: "company.employee_count"},
          open(os.path.join(T, "iplan.json"), "w"))
imp = run("import", old, "--plan", os.path.join(T, "iplan.json"))
assert {s["key"] for s in imp["saved"]} == {"company.year_established", "company.employee_count"}
keys = run("keys")
assert keys["company.employee_count"]["on_file"] and keys["tax.ein"]["on_file"] and "value" not in json.dumps(keys)

# ---- documents: encrypted in, decrypted only outside the repo, rules apply, rotation keeps them readable ----
w9 = os.path.join(T, "W-9.pdf"); open(w9, "wb").write(b"%PDF-1.4 fake W-9 " + EIN.encode())
coi = os.path.join(T, "COI.pdf"); open(coi, "wb").write(b"%PDF-1.4 fake insurance certificate")
run("store-doc", "w9", w9); run("store-doc", "certificate_of_insurance", coi)
assert EIN.encode() not in open(os.path.join(CONFIG, "docs", "w9.age"), "rb").read()
docs = run("docs"); assert docs["w9"]["release"] == "auto" and docs["certificate_of_insurance"]["release"] == "approve"
outside = tempfile.mkdtemp()
got = run("open-doc", "w9", outside, "--recipient", "ap@acme.example")["path"]
assert open(got, "rb").read() == open(w9, "rb").read()
p = subprocess.run([sys.executable, os.path.join(HERE, "formfill.py"), "--config", CONFIG, "--ops", OPS,
                    "open-doc", "certificate_of_insurance", outside], capture_output=True, text=True)
assert p.returncode != 0 and "approve" in p.stderr               # ask-each-time documents need --approved
p = subprocess.run([sys.executable, os.path.join(HERE, "formfill.py"), "--config", CONFIG, "--ops", OPS,
                    "open-doc", "w9", os.path.join(T, "inside")], capture_output=True, text=True, cwd=T)
assert p.returncode != 0                                         # never decrypt inside the repo
run("rotate")
assert os.path.exists(os.environ["OFFICER_CAN_KEY"] + ".old")
assert open(run("open-doc", "w9", tempfile.mkdtemp())["path"], "rb").read() == open(w9, "rb").read()
assert run("keys")["tax.ein"]["on_file"]
r = run("fill", form1, out1, "--plan", os.path.join(T, "plan.json"))
assert {f["label"]: f["status"] for f in r["fields"]}["Federal Tax ID (EIN)"] == "filled"

shutil.rmtree(T)
print("formfill: all checks passed")
